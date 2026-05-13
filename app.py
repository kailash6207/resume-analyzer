from flask import Flask, render_template, request
import os
import pdfplumber
from docx import Document

# Import your new utility modules
from utils.parser import extract_email, extract_phone, extract_skills
from utils.matcher import match_skills, calculate_score

app = Flask(__name__)

UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

def extract_pdf_text(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    return text

# Function to extract text from DOCX (Updated to read tables!)
def extract_docx_text(file_path):
    doc = Document(file_path)
    text = ""
    
    # 1. Extract regular paragraphs
    for para in doc.paragraphs:
        text += para.text + "\n"
        
    # 2. Extract text from inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + " "
        text += "\n"
        
    return text

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/analyze', methods=['POST'])
def analyze():
    if 'resume' not in request.files:
        return "No file uploaded"

    file = request.files['resume']
    
    # 1. Grab the Job Description text from the new HTML form
    job_description = request.form.get('job_description', '') 
    
    print(f"DEBUG: The Job Description received is: '{job_description}'")

    if file.filename == '':
        return "No selected file"

    file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
    file.save(file_path)

    extracted_text = ""

    if file.filename.endswith('.pdf'):
        extracted_text = extract_pdf_text(file_path)
    elif file.filename.endswith('.docx'):
        extracted_text = extract_docx_text(file_path)
    else:
        return "Unsupported file format. Upload PDF or DOCX."

    # 2. Parse the Resume
    email = extract_email(extracted_text)
    phone = extract_phone(extracted_text)
    resume_skills = extract_skills(extracted_text)
    
    # 3. Parse the Job Description (if the user pasted one)
    target_skills = None
    if job_description.strip():
        target_skills = [skill.strip() for skill in job_description.split(',')]
    
    # 4. Match the resume skills against the target JD skills
    match_data = match_skills(resume_skills, target_skills)
    extracted_data = {'email': email, 'phone': phone, 'skills': resume_skills}
    score = calculate_score(extracted_data, match_data)

    return render_template('result.html', 
                           email=email, 
                           phone=phone, 
                           skills=resume_skills, 
                           score=score,
                           matched=match_data['matched'])

if __name__ == '__main__':
    app.run(debug=True)